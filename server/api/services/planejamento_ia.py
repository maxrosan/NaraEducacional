"""
Serviço de IA do Planejamento.

Funções puras chamadas pelas views REST:
- extrair_texto_arquivo(uploaded_file) -> str
- salvar_arquivo_planejamento(uploaded_file, turma_id, dia_semana) -> dict
- sugerir_atividades_a_partir_de_arquivo(texto) -> str
- sugerir_atividades_a_partir_de_prompt(prompt, contexto_turma, cliente_id) -> str
- sugerir_habilidades_bncc(atividades_texto, ano_serie, limite, cliente_id) -> dict
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import re
from datetime import date
from pathlib import Path
from typing import Optional
from uuid import uuid4

from api.models import HabilidadeBNCC
from api.ia_utils import _get_int_env
from api.openai_client import get_openai_client
from api.storage import (
    delete_from_s3,
    generate_presigned_url,
    is_s3_configured,
    upload_bytes_to_storage,
)
from api.services.prompt_resolver import resolver_prompt


logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Prompts de fallback (usados quando o banco não tem registro para a categoria)
# ---------------------------------------------------------------------------

_PROMPT_PLANEJAMENTO_ATIVIDADES_FALLBACK = (
    "Você é uma assistente pedagógica especializada em Educação Infantil e "
    "anos iniciais do Ensino Fundamental no Brasil. A professora descreverá "
    "brevemente o objetivo, tema ou atividade que quer trabalhar. Sua tarefa "
    "é propor uma sequência curta e prática de atividades para um dia de "
    "aula, em português do Brasil, em formato de bullets, sem títulos "
    "extras. Não cite habilidades BNCC nesta resposta."
)

_PROMPT_PLANEJAMENTO_BNCC_FALLBACK = (
    "Você é uma especialista pedagógica. Selecione as habilidades BNCC mais "
    "relevantes para as atividades descritas. USE APENAS habilidades "
    "presentes na lista enviada — não invente códigos. "
    "Responda em JSON no formato: "
    '{"habilidades":[{"id":"...","codigo":"...","justificativa":"..."}]}'
)


# ---------------------------------------------------------------------------
# Constantes de validação de arquivo
# ---------------------------------------------------------------------------

ALLOWED_PLANEJAMENTO_MIME_TYPES = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
ALLOWED_PLANEJAMENTO_EXTENSIONS = {".pdf", ".doc", ".docx"}


def _dir_upload_local() -> Path:
    """Pasta do fallback local (sem S3). Upload, URL e remoção usam a MESMA —
    antes a remoção ignorava PLANEJAMENTO_UPLOAD_DIR e apagava em "uploads/"."""
    return Path(os.getenv("PLANEJAMENTO_UPLOAD_DIR", "uploads"))


MAX_PLANEJAMENTO_SIZE_BYTES = (
    _get_int_env("MAX_PLANEJAMENTO_UPLOAD_MB", 10) * 1024 * 1024
)


# ---------------------------------------------------------------------------
# Utilidades internas
# ---------------------------------------------------------------------------


def _normalizar_extensao(uploaded_file) -> str:
    content_type = (uploaded_file.content_type or "").lower().split(";")[0].strip()
    original_ext = os.path.splitext(uploaded_file.name or "")[1].lower()

    mime_to_ext = {
        "application/pdf": ".pdf",
        "application/msword": ".doc",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    }

    if content_type in mime_to_ext:
        return mime_to_ext[content_type]

    if original_ext in ALLOWED_PLANEJAMENTO_EXTENSIONS:
        return original_ext

    return ""


def _ler_bytes(uploaded_file) -> bytes:
    try:
        uploaded_file.seek(0)
    except Exception:
        pass
    content = uploaded_file.read()
    try:
        uploaded_file.seek(0)
    except Exception:
        pass
    return content


def _extrair_json_de_resposta_ia(conteudo: str) -> dict:
    if not conteudo:
        return {}

    texto = conteudo.strip()

    if texto.startswith("```"):
        texto = re.sub(r"^```(?:json)?\s*", "", texto, flags=re.IGNORECASE)
        texto = re.sub(r"\s*```$", "", texto)

    if "{" in texto and "}" in texto:
        inicio = texto.find("{")
        fim = texto.rfind("}")
        if inicio != -1 and fim != -1 and fim > inicio:
            texto = texto[inicio:fim + 1]

    return json.loads(texto) if texto else {}


def _modelo_ia() -> str:
    return os.getenv("OPENAI_PLANNING_MODEL", "gpt-4o-mini")


# ---------------------------------------------------------------------------
# Extração de texto
# ---------------------------------------------------------------------------


def extrair_texto_arquivo(uploaded_file) -> str:
    extensao = _normalizar_extensao(uploaded_file)
    if extensao not in ALLOWED_PLANEJAMENTO_EXTENSIONS:
        raise ValueError(
            "Formato de arquivo não suportado. Envie um PDF ou DOC/DOCX."
        )

    conteudo = _ler_bytes(uploaded_file)
    if not conteudo:
        raise ValueError("Arquivo vazio.")

    try:
        if extensao == ".pdf":
            return _extrair_texto_pdf(conteudo)
        if extensao in {".doc", ".docx"}:
            return _extrair_texto_docx(conteudo)
    except ValueError:
        raise
    except Exception as exc:
        logger.warning("Falha ao extrair texto do arquivo: %s", exc)
        raise ValueError(
            "Não foi possível ler o conteúdo do arquivo enviado. "
            "Confirme que ele não está protegido por senha ou corrompido."
        ) from exc

    return ""


def _extrair_texto_pdf(conteudo: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(conteudo))
    paginas = []
    for page in reader.pages:
        try:
            paginas.append(page.extract_text() or "")
        except Exception:
            paginas.append("")
    texto = "\n".join(p.strip() for p in paginas if p.strip())
    if not texto.strip():
        raise ValueError(
            "PDF sem texto extraível (provavelmente é uma imagem digitalizada). "
            "Envie um arquivo com texto pesquisável."
        )
    return texto


def _extrair_texto_docx(conteudo: bytes) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise ValueError(
            "Biblioteca de leitura de DOCX indisponível no servidor."
        ) from exc

    try:
        documento = docx.Document(io.BytesIO(conteudo))
    except Exception as exc:
        raise ValueError(
            "Não foi possível ler o documento. Salve como .docx (Word moderno) "
            "e tente novamente."
        ) from exc

    paragrafos = [p.text.strip() for p in documento.paragraphs if p.text.strip()]
    if not paragrafos:
        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    txt = celula.text.strip()
                    if txt:
                        paragrafos.append(txt)

    texto = "\n".join(paragrafos)
    if not texto.strip():
        raise ValueError("Documento sem conteúdo de texto identificável.")
    return texto


# ---------------------------------------------------------------------------
# Persistência do arquivo (S3 + fallback local em dev)
# ---------------------------------------------------------------------------


def salvar_arquivo_planejamento(
    uploaded_file,
    turma_id: str,
    dia_semana: str,
) -> dict:
    extensao = _normalizar_extensao(uploaded_file) or ".bin"
    conteudo = _ler_bytes(uploaded_file)

    digest = hashlib.sha256(conteudo).hexdigest()[:16]
    safe_turma = re.sub(r"[^A-Za-z0-9_-]", "_", str(turma_id or "sem-turma"))[:48]
    safe_dia = re.sub(r"[^A-Za-z0-9_-]", "_", str(dia_semana or "sem-dia"))[:16]
    key = f"planejamentos/{safe_turma}/{safe_dia}/{digest}-{uuid4().hex[:8]}{extensao}"

    nome_original = uploaded_file.name or f"planejamento{extensao}"
    content_type = (uploaded_file.content_type or "application/octet-stream")

    if is_s3_configured():
        storage_key, arquivo_url = upload_bytes_to_storage(
            key, conteudo, content_type=content_type
        )
    else:
        storage_key = key
        base_dir = _dir_upload_local()
        destino = base_dir / key
        destino.parent.mkdir(parents=True, exist_ok=True)
        with destino.open("wb") as fh:
            fh.write(conteudo)
        arquivo_url = f"/{destino.as_posix()}"
        logger.info(
            "[planejamento_ia] S3 não configurado; arquivo gravado em %s", destino
        )

    return {
        "storage_key": storage_key,
        "arquivo_url": arquivo_url,
        "arquivo_nome_original": nome_original,
        "arquivo_content_type": content_type,
    }


def regenerar_url_arquivo(storage_key: Optional[str]) -> Optional[str]:
    if not storage_key:
        return None
    if not is_s3_configured():
        return f"/{(_dir_upload_local() / storage_key).as_posix()}"
    return generate_presigned_url(storage_key)


def remover_arquivo_planejamento(storage_key: Optional[str]) -> None:
    if not storage_key:
        return
    if is_s3_configured():
        delete_from_s3(storage_key)
        return
    try:
        caminho = _dir_upload_local() / storage_key
        if caminho.exists():
            caminho.unlink()
    except Exception as exc:
        logger.warning("Falha ao remover arquivo local %s: %s", storage_key, exc)


# ---------------------------------------------------------------------------
# Sugestões de atividades (LLM)
# ---------------------------------------------------------------------------


def _chamar_openai_text(system_prompt: str, user_prompt: str, *, max_tokens: int = 800) -> str:
    client = get_openai_client()
    response = client.chat.completions.create(
        model=_modelo_ia(),
        temperature=0.3,
        max_tokens=max_tokens,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )
    return (response.choices[0].message.content or "").strip()


def extrair_atividades_por_dia(texto_extraido: str) -> dict:
    """
    Lê o texto bruto extraído de um arquivo de planejamento e devolve,
    em UMA chamada à OpenAI, atividades estruturadas por data.
    Prompt técnico com JSON schema — não editável pelo usuário.
    """
    vazio = {
        "dias": [],
        "data_inicio": None,
        "data_fim": None,
        "evidencia": "",
        "confianca": None,
        "fallback_texto_unico": "",
    }

    texto = (texto_extraido or "").strip()
    if not texto:
        return vazio
    if not _texto_tem_indicio_de_data(texto):
        return vazio

    # O gpt-4o-mini tem janela de 128k tokens, então um planejamento semanal
    # (tipicamente ~15-30k chars) cabe folgado. O limite antigo (12k) era baixo
    # demais e, pior, truncava por "cabeça + cauda" — o miolo do documento, que
    # num plano de seg a sex contém justamente a quarta/quinta, era descartado
    # ANTES de chegar à IA, fazendo dias sumirem. Mandamos o texto inteiro até um
    # Detecta deterministicamente, ANTES de truncar, as datas de dias úteis
    # presentes no texto. Servem de âncora no prompt: a IA não precisa "lembrar"
    # de cada dia — recebe a lista exata de datas a preencher. Sem isso, o
    # gpt-4o-mini às vezes deixava quinta/sexta sem data (e o frontend descarta
    # itens sem data, sumindo com o dia).
    datas_uteis = _datas_uteis_no_texto(texto)

    # teto bem mais alto; só documentos realmente gigantes sofrem truncamento.
    limite = _get_int_env("OPENAI_PLANNING_MAX_CHARS", 48000)
    if len(texto) > limite:
        meio = limite // 2
        texto = texto[:meio] + "\n\n[...]\n\n" + texto[-meio:]

    hoje = date.today().isoformat()
    system_prompt = (
        "Você é uma assistente pedagógica brasileira. Receba o texto bruto "
        "de um planejamento e devolva ESTRITAMENTE em JSON, com este "
        "esquema:\n"
        '{"dias":[{"data":"YYYY-MM-DD"|null,'
        '"dia_semana":"segunda"|"terca"|"quarta"|"quinta"|"sexta",'
        '"atividades":"texto em prosa curta ou bullets simples"}],'
        '"data_inicio":"YYYY-MM-DD"|null,'
        '"data_fim":"YYYY-MM-DD"|null,'
        '"evidencia":"trecho onde a data/intervalo apareceu",'
        '"confianca":"alta"|"media"|"baixa",'
        '"fallback_texto_unico":"texto consolidado quando NÃO der pra '
        'separar por dia"}\n'
        "Regras:\n"
        "- Devolva UM item por dia útil COM CONTEÚDO no texto (segunda a "
        "sexta). Se o documento cobre a semana inteira, devolva os 5 dias.\n"
        "- SEMPRE preencha o campo \"data\" (YYYY-MM-DD) quando a data do dia "
        "existir no texto. NÃO deixe \"data\" em branco se der para saber a "
        "data — derive o \"dia_semana\" a partir da \"data\".\n"
        "- Cada par (\"data\"/\"dia_semana\") aparece UMA ÚNICA vez. NUNCA "
        "repita o mesmo dia em dois itens, nem rotule dois dias diferentes com "
        "o mesmo \"dia_semana\".\n"
        "- Cada bloco \"atividades\" deve ser em português do Brasil, em "
        "bullets simples começando com \"- \". CADA atividade fica em sua "
        "própria linha, separadas por quebra de linha (\\n). NUNCA junte "
        "vários itens numa mesma linha.\n"
        "- Use APENAS o que está no texto; não invente atividades nem datas.\n"
        "- Sábados e domingos devem ser ignorados.\n"
        "- Se o texto tiver apenas dia/mês sem ano, escolha o ano que torne a "
        "data mais próxima de hoje.\n"
        "- \"data_inicio\" e \"data_fim\" cobrem o intervalo geral coberto "
        "pelo arquivo. Se não houver datas, use null.\n"
        "- Se NÃO conseguir separar por dia, deixe \"dias\" como [] e "
        "preencha \"fallback_texto_unico\" também com bullets em linhas "
        "separadas por \\n."
    )
    if datas_uteis:
        ancora_datas = (
            "\n\nDatas de dias ÚTEIS já identificadas no documento (formato "
            "ISO). Devolva UMA entrada para CADA UMA que tiver conteúdo no "
            "texto, usando exatamente este valor no campo \"data\":\n- "
            + "\n- ".join(datas_uteis)
        )
    else:
        ancora_datas = ""
    user_prompt = (
        f"Data de hoje: {hoje}.\n\n"
        f"Texto bruto extraído do arquivo:\n{texto}"
        f"{ancora_datas}"
    )

    try:
        client = get_openai_client()
        response = client.chat.completions.create(
            model=_modelo_ia(),
            temperature=0.2,
            # 5 dias úteis com bullets detalhados podem passar de 2400 tokens;
            # se a resposta JSON for cortada por limite, o parse falha e a
            # extração inteira cai para vazio. Folga maior evita perder dias por
            # JSON truncado (mesmo motivo do ajuste feito em audio.py).
            max_tokens=_get_int_env("OPENAI_PLANNING_MAX_TOKENS", 4000),
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        content = (response.choices[0].message.content or "").strip()
        parsed = _extrair_json_de_resposta_ia(content)
    except Exception as exc:
        logger.warning("Falha ao extrair atividades por dia: %s", exc)
        return vazio

    resultado = _normalizar_extracao_atividades(parsed)
    if not resultado["dias"] and not resultado["fallback_texto_unico"]:
        # A IA respondeu, mas nada sobreviveu à normalização. Logamos para
        # diferenciar "sem sugestões por falha/silêncio da IA" de "arquivo sem
        # datas" — antes os dois casos eram indistinguíveis no servidor.
        logger.info(
            "Extração sem atividades aproveitáveis (texto=%d chars). "
            "Resposta crua da IA: %.300s",
            len(texto),
            content,
        )
    return resultado


def _normalizar_extracao_atividades(parsed: dict) -> dict:
    dias_raw = parsed.get("dias") or []
    if not isinstance(dias_raw, list):
        dias_raw = []

    dias_normalizados: list[dict] = []
    vistos_dias: set[str] = set()
    for item in dias_raw:
        if not isinstance(item, dict):
            continue
        atividades = _normalizar_bullets(item.get("atividades") or "")
        if not atividades:
            continue

        data_iso = (item.get("data") or "").strip() or None
        if data_iso:
            data_iso = _ajustar_ano_proximo(data_iso)
        dia_semana = (item.get("dia_semana") or "").strip().lower()
        if dia_semana == "terça":
            dia_semana = "terca"

        if data_iso:
            try:
                data_obj = date.fromisoformat(data_iso)
            except (ValueError, TypeError):
                data_obj = None
            if data_obj is not None:
                if data_obj.weekday() > 4:
                    # Data caiu em sábado/domingo — quase sempre efeito de
                    # inferência de ano equivocada (ex.: arquivo intitulado
                    # "2025" cujo calendário, na verdade, é de outro ano, então
                    # 18/05 vira domingo). Se a IA mandou um rótulo de dia útil
                    # válido, preservamos a atividade sob esse rótulo e apenas
                    # descartamos a data não confiável, em vez de jogar o dia
                    # inteiro fora (era a causa de dias sumirem silenciosamente).
                    if dia_semana in _DIAS_SEMANA_VALIDOS:
                        data_iso = None
                    else:
                        continue
                else:
                    # Dia útil com data confiável: a DATA manda. Derivamos o
                    # dia_semana dela SEMPRE, sobrescrevendo o rótulo que a IA
                    # eventualmente mandou errado (ex.: 18/05 = segunda chegando
                    # rotulado como "quinta"). Sem isso, o frontend agrupava e
                    # mapeava o conteúdo no dia errado na importação — segunda
                    # caía na quinta, dois dias colidiam no mesmo rótulo e parte
                    # das atividades se perdia.
                    dia_semana = _DIA_POR_WEEKDAY[data_obj.weekday()]
        else:
            if dia_semana not in _DIAS_SEMANA_VALIDOS:
                continue

        chave = f"{data_iso or '-'}|{dia_semana}"
        if chave in vistos_dias:
            continue
        vistos_dias.add(chave)

        dias_normalizados.append({
            "data": data_iso,
            "dia_semana": dia_semana,
            "atividades": atividades,
        })

    data_inicio = (parsed.get("data_inicio") or "").strip() or None
    data_fim = (parsed.get("data_fim") or "").strip() or None
    if data_inicio:
        data_inicio = _ajustar_ano_proximo(data_inicio)
    if data_fim:
        data_fim = _ajustar_ano_proximo(data_fim)

    if dias_normalizados:
        datas_dias = [d["data"] for d in dias_normalizados if d["data"]]
        if datas_dias:
            if not data_inicio:
                data_inicio = min(datas_dias)
            if not data_fim:
                data_fim = max(datas_dias)

    confianca = (parsed.get("confianca") or "").strip().lower()
    if confianca not in {"alta", "media", "baixa"}:
        confianca = "media" if dias_normalizados or data_inicio else "baixa"

    evidencia = (parsed.get("evidencia") or "").strip()
    if len(evidencia) > 240:
        evidencia = evidencia[:240]

    fallback_texto = _normalizar_bullets(parsed.get("fallback_texto_unico") or "")

    return {
        "dias": dias_normalizados,
        "data_inicio": data_inicio,
        "data_fim": data_fim,
        "evidencia": evidencia,
        "confianca": confianca,
        "fallback_texto_unico": fallback_texto,
    }


def sugerir_atividades_a_partir_de_arquivo(texto_extraido: str) -> str:
    texto = (texto_extraido or "").strip()
    if not texto:
        raise ValueError("Texto extraído vazio; nada a sugerir.")

    extracao = extrair_atividades_por_dia(texto_extraido)
    if extracao["dias"]:
        partes = []
        for d in extracao["dias"]:
            label = (d["dia_semana"] or "").capitalize()
            if d["data"]:
                label = f"{label} ({d['data']})" if label else d["data"]
            cabecalho = f"{label}:" if label else ""
            partes.append(
                f"{cabecalho}\n{d['atividades']}".strip()
            )
        return "\n\n".join(partes)
    return extracao.get("fallback_texto_unico") or ""


# ---------------------------------------------------------------------------
# Detecção da data/intervalo do planejamento
# ---------------------------------------------------------------------------


# Tolera espaços/tabs ao redor do separador: muitos modelos de planejamento
# escrevem a data como "18 / 05" (com espaços). Sem isso o gate
# ``_texto_tem_indicio_de_data`` não enxergava a data e, em arquivos sem os
# nomes dos dias escritos por extenso, devolvia "nenhuma sugestão" sem nem
# chamar a IA. Limitamos a [ \t] (não \s) para não casar através de quebras de
# linha e juntar números de linhas diferentes.
#
# As âncoras usam lookarounds de dígito — (?<!\d) / (?!\d) — em vez de \b. O
# \b falhava quando a data vinha colada a uma letra, que é exatamente como o
# pypdf extrai os planejamentos: "DATA: 19 / 05PROFESSORA..." (sem espaço entre
# "05" e "PR"). Como "5" e "P" são ambos word-chars, não há \b ali e a data
# era ignorada — só as datas seguidas de espaço (cabeçalho "SEMANA: 18/05 A
# 22/05") eram detectadas. Com (?!\d) basta que não venha outro dígito.
_DATE_HINT_REGEX = re.compile(
    r"(?<!\d)(\d{1,2})[ \t]*[/\-\.][ \t]*(\d{1,2})"
    r"(?:[ \t]*[/\-\.][ \t]*(\d{2,4}))?(?!\d)"
)
_DAY_NAME_HINT_REGEX = re.compile(
    r"\b(segunda|ter[cç]a|quarta|quinta|sexta)(?:[\s\-]?feira)?\b",
    re.IGNORECASE,
)


def _texto_tem_indicio_de_data(texto: str) -> bool:
    if not texto:
        return False
    if _DATE_HINT_REGEX.search(texto):
        return True
    return bool(_DAY_NAME_HINT_REGEX.search(texto))


_DIAS_SEMANA_VALIDOS = {"segunda", "terca", "quarta", "quinta", "sexta"}
_DIA_POR_WEEKDAY = {0: "segunda", 1: "terca", 2: "quarta", 3: "quinta", 4: "sexta"}


_BULLET_MARKER_REGEX = re.compile(r"(?:- |• |\* |\d+\.\s|\d+\)\s)")
_BULLET_SPLIT_REGEX = re.compile(
    r"(?<!^)(?<!\n)\s+(?=(?:- |• |\* |\d+\.\s|\d+\)\s))"
)


def _normalizar_bullets(texto: str) -> str:
    if not texto:
        return ""

    bruto = texto.strip()
    if "\n" in bruto:
        linhas = [linha.strip() for linha in bruto.split("\n")]
        return "\n".join(linha for linha in linhas if linha)

    if len(_BULLET_MARKER_REGEX.findall(bruto)) < 2:
        return bruto

    quebrado = _BULLET_SPLIT_REGEX.sub("\n", bruto)
    linhas = [linha.strip() for linha in quebrado.split("\n")]
    return "\n".join(linha for linha in linhas if linha)


def _ajustar_ano_proximo(data_iso: str) -> Optional[str]:
    if not data_iso:
        return None
    try:
        ano, mes, dia = (int(p) for p in data_iso.split("-"))
        candidata = date(ano, mes, dia)
    except (ValueError, AttributeError):
        return None

    hoje = date.today()
    if abs((candidata - hoje).days) <= 180:
        return candidata.isoformat()

    candidatas = []
    for ano_tentativa in (hoje.year - 1, hoje.year, hoje.year + 1):
        try:
            candidatas.append(date(ano_tentativa, mes, dia))
        except ValueError:
            continue
    if not candidatas:
        return candidata.isoformat()
    melhor = min(candidatas, key=lambda d: abs((d - hoje).days))
    return melhor.isoformat()


def _datas_uteis_no_texto(texto: str) -> list[str]:
    """
    Encontra, de forma determinística, as datas de dias ÚTEIS presentes no
    texto (ex.: "DATA: 18 / 05", "SEMANA: 18/05 A 22/05") e as devolve em ISO,
    sem repetição e ordenadas. O ano ausente é resolvido por proximidade de hoje
    (mesma lógica de `_ajustar_ano_proximo`). Sábados/domingos são descartados.

    Serve de âncora no prompt de extração: a IA recebe a lista exata de datas a
    preencher, em vez de depender de "lembrar" de cada dia.
    """
    if not texto:
        return []

    vistas: list[str] = []
    seen: set[str] = set()
    for m in _DATE_HINT_REGEX.finditer(texto):
        try:
            dia = int(m.group(1))
            mes = int(m.group(2))
        except (TypeError, ValueError):
            continue
        if not (1 <= dia <= 31 and 1 <= mes <= 12):
            continue

        ano_grp = m.group(3)
        if ano_grp:
            ano = int(ano_grp)
            if ano < 100:
                ano += 2000
        else:
            ano = date.today().year

        try:
            iso = _ajustar_ano_proximo(f"{ano:04d}-{mes:02d}-{dia:02d}")
        except ValueError:
            continue
        if not iso:
            continue
        try:
            d = date.fromisoformat(iso)
        except ValueError:
            continue
        if d.weekday() > 4:  # ignora fim de semana
            continue
        if iso not in seen:
            seen.add(iso)
            vistas.append(iso)

    vistas.sort()
    return vistas


def detectar_data_planejamento(texto_extraido: str) -> Optional[dict]:
    extracao = extrair_atividades_por_dia(texto_extraido)
    if not extracao.get("data_inicio"):
        return None
    return {
        "data_inicio": extracao["data_inicio"],
        "data_fim": extracao["data_fim"],
        "evidencia": extracao["evidencia"],
        "confianca": extracao["confianca"] or "media",
    }


def sugerir_atividades_a_partir_de_prompt(
    prompt: str,
    contexto_turma: str = "",
    cliente_id: str = None,
) -> str:
    """
    Assistente IA livre — gera sugestão a partir de descrição da professora.
    O system_prompt é resolvido do banco (categoria "Planejamento"),
    com fallback para _PROMPT_PLANEJAMENTO_ATIVIDADES_FALLBACK.
    """
    prompt_limpo = (prompt or "").strip()
    if len(prompt_limpo) < 10:
        raise ValueError("Descreva o objetivo/tema com pelo menos 10 caracteres.")

    system_prompt = (
        resolver_prompt("Planejamento", cliente_id=cliente_id)
        or _PROMPT_PLANEJAMENTO_ATIVIDADES_FALLBACK
    )

    contexto = (contexto_turma or "").strip()
    user_prompt = (
        (f"Contexto da turma: {contexto}\n\n" if contexto else "")
        + f"Descrição da professora:\n{prompt_limpo}"
    )
    return _chamar_openai_text(system_prompt, user_prompt)


# ---------------------------------------------------------------------------
# Sugestão de Habilidades BNCC
# ---------------------------------------------------------------------------


def _candidatos_bncc(ano_serie: str = "") -> list[dict]:
    """
    No schema novo, HabilidadeBNCC já É o catálogo oficial (Pergunta.habilidade_bncc
    é uma FK de verdade pra cá, não um código texto solto) — não precisa mais
    derivar candidatos a partir de perguntas.
    """
    queryset = HabilidadeBNCC.objects.filter(ativa=True)
    if ano_serie:
        queryset = queryset.filter(ano_serie__iexact=ano_serie)

    candidatos = [
        {
            "id": str(h.id),
            "codigo": h.codigo,
            "descricao": h.descricao,
            "componente_curricular": h.componente_curricular or "",
            "ano_serie": h.ano_serie or "",
            "campo_atuacao": h.campo_atuacao or "",
        }
        for h in queryset
    ]
    return candidatos


def _fallback_por_palavras(
    candidatos: list[dict], atividades_texto: str, limite: int
) -> list[dict]:
    stopwords = {
        "para", "com", "sobre", "entre", "mais", "menos", "onde", "quando",
        "como", "isso", "essa", "este", "esta", "esse", "pela", "pelo",
        "dias", "semana", "turma", "atividade", "atividades", "planejamento",
    }
    tokens = [
        token
        for token in re.split(r"[^0-9a-zA-ZÀ-ÿ]+", atividades_texto.lower())
        if len(token) >= 4 and token not in stopwords
    ]
    tokens = list(dict.fromkeys(tokens))[:12]

    def score(habilidade):
        haystack = (
            f"{habilidade['codigo']} {habilidade['descricao']} "
            f"{habilidade['componente_curricular']} {habilidade['ano_serie']} "
            f"{habilidade['campo_atuacao']}"
        ).lower()
        return sum(1 for token in tokens if token in haystack)

    ordenadas = sorted(candidatos, key=score, reverse=True)
    selecionadas = ordenadas[: max(1, limite)]
    return [
        {
            **h,
            "justificativa": "Sugestão baseada em correspondência por palavras-chave.",
        }
        for h in selecionadas
    ]


def sugerir_habilidades_bncc(
    atividades_texto: str,
    ano_serie: str = "",
    limite: int = 6,
    cliente_id: str = None,
) -> dict:
    """
    Devolve {habilidades, origem, mensagem?}.
    O system_prompt é resolvido do banco (categoria "Planejamento"),
    com fallback para _PROMPT_PLANEJAMENTO_BNCC_FALLBACK.
    """
    texto = (atividades_texto or "").strip()
    if len(texto) < 10:
        raise ValueError(
            "Atividades muito curtas para sugerir habilidades. "
            "Descreva ou cole as atividades antes de gerar a BNCC."
        )

    limite = max(1, min(int(limite or 6), 10))
    candidatos = _candidatos_bncc(ano_serie=ano_serie)
    if not candidatos:
        return {
            "habilidades": [],
            "origem": "fallback",
            "mensagem": "Nenhuma habilidade ativa cadastrada no catálogo BNCC.",
        }

    fallback_resultado = _fallback_por_palavras(candidatos, texto, limite)

    payload_candidatos = [
        {
            "id": h["id"],
            "codigo": h["codigo"],
            "descricao": h["descricao"],
            "componente_curricular": h["componente_curricular"],
            "ano_serie": h["ano_serie"],
            "campo_atuacao": h["campo_atuacao"],
        }
        for h in candidatos[:120]
    ]

    system_prompt = (
        resolver_prompt("Planejamento", cliente_id=cliente_id)
        or _PROMPT_PLANEJAMENTO_BNCC_FALLBACK
    )

    user_prompt = (
        f"Ano/série: {ano_serie or 'não informado'}\n\n"
        f"Atividades propostas:\n{texto}\n\n"
        f"Habilidades candidatas (use apenas estas):\n"
        f"{json.dumps(payload_candidatos, ensure_ascii=False)}"
    )

    try:
        client = get_openai_client()
        response = client.chat.completions.create(
            model=_modelo_ia(),
            temperature=0.2,
            max_tokens=600,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
        )
        content = (response.choices[0].message.content or "").strip()
        parsed = _extrair_json_de_resposta_ia(content)
        habilidades_ia = parsed.get("habilidades", [])
        if not isinstance(habilidades_ia, list):
            raise ValueError("Resposta da IA fora do formato esperado.")

        candidatos_por_id = {h["id"]: h for h in payload_candidatos}
        candidatos_por_codigo = {h["codigo"]: h for h in payload_candidatos}

        selecionadas: list[dict] = []
        vistos: set[str] = set()
        for habilidade in habilidades_ia:
            habilidade_id = str(habilidade.get("id", "")).strip()
            habilidade_codigo = str(habilidade.get("codigo", "")).strip()
            base = candidatos_por_id.get(habilidade_id) or candidatos_por_codigo.get(
                habilidade_codigo
            )
            if not base or base["codigo"] in vistos:
                continue
            justificativa = (habilidade.get("justificativa") or "").strip()
            selecionadas.append(
                {
                    **base,
                    "justificativa": justificativa or "Sugestão gerada pela IA.",
                }
            )
            vistos.add(base["codigo"])
            if len(selecionadas) >= limite:
                break

        if not selecionadas:
            logger.warning("IA não retornou habilidades válidas; usando fallback.")
            return {
                "habilidades": fallback_resultado,
                "origem": "fallback",
                "mensagem": "IA não retornou sugestões válidas; usando fallback.",
            }

        return {"habilidades": selecionadas, "origem": "ia"}

    except Exception as exc:
        logger.warning("Fallback acionado na sugestão BNCC: %s", exc)
        return {
            "habilidades": fallback_resultado,
            "origem": "fallback",
            "mensagem": "IA indisponível; sugestões geradas por palavras-chave.",
        }